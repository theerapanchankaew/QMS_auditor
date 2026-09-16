from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
import hashlib

from .models import SourceDocument, SourceKind, SourceSpan


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


@dataclass
class ParseResult:
    document: SourceDocument
    spans: list[SourceSpan]


class DocumentParser:
    kind: SourceKind

    def parse(self, path: str | Path, *, source_id: str, organization_id: str,
              audit_case_id: str | None = None, controlled: bool = True,
              version: str | None = None, authority: str | None = None,
              ingested_at=None) -> ParseResult:
        raise NotImplementedError


class TextParser(DocumentParser):
    kind = SourceKind.TEXT

    def parse(self, path, *, source_id, organization_id, audit_case_id=None, controlled=True,
              version=None, authority=None, ingested_at=None) -> ParseResult:
        from datetime import datetime, timezone
        p = Path(path)
        raw = p.read_bytes()
        text = raw.decode("utf-8", errors="replace")
        doc = SourceDocument(
            source_id=source_id, organization_id=organization_id, audit_case_id=audit_case_id,
            title=p.name, kind=self.kind, uri=str(p), version=version, authority=authority,
            controlled=controlled, sha256=sha256_bytes(raw), size_bytes=len(raw),
            ingested_at=ingested_at or datetime.now(timezone.utc),
        )
        spans = []
        offset = 0
        for i, para in enumerate([x.strip() for x in text.splitlines() if x.strip()]):
            start = text.find(para, offset)
            end = start + len(para)
            spans.append(SourceSpan(
                span_id=f"{source_id}:P{i}", source_id=source_id, paragraph_index=i,
                char_start=start, char_end=end, text=para, span_hash=sha256_text(para)
            ))
            offset = end
        return ParseResult(doc, spans)


class PDFParser(DocumentParser):
    kind = SourceKind.PDF

    def parse(self, path, *, source_id, organization_id, audit_case_id=None, controlled=True,
              version=None, authority=None, ingested_at=None) -> ParseResult:
        from datetime import datetime, timezone
        import fitz
        p = Path(path)
        raw = p.read_bytes()
        doc_meta = SourceDocument(
            source_id=source_id, organization_id=organization_id, audit_case_id=audit_case_id,
            title=p.name, kind=self.kind, uri=str(p), version=version, authority=authority,
            controlled=controlled, sha256=sha256_bytes(raw), size_bytes=len(raw),
            ingested_at=ingested_at or datetime.now(timezone.utc),
        )
        spans: list[SourceSpan] = []
        with fitz.open(stream=raw, filetype="pdf") as pdf:
            for page_no, page in enumerate(pdf, start=1):
                blocks = page.get_text("blocks")
                for bi, block in enumerate(blocks):
                    text = str(block[4]).strip()
                    if not text:
                        continue
                    sid = f"{source_id}:PG{page_no}:B{bi}"
                    spans.append(SourceSpan(
                        span_id=sid, source_id=source_id, page=page_no, text=text,
                        span_hash=sha256_text(text), metadata={"bbox": list(block[:4])}
                    ))
        return ParseResult(doc_meta, spans)


class DOCXParser(DocumentParser):
    kind = SourceKind.DOCX

    def parse(self, path, *, source_id, organization_id, audit_case_id=None, controlled=True,
              version=None, authority=None, ingested_at=None) -> ParseResult:
        from datetime import datetime, timezone
        from docx import Document
        p = Path(path)
        raw = p.read_bytes()
        doc_meta = SourceDocument(
            source_id=source_id, organization_id=organization_id, audit_case_id=audit_case_id,
            title=p.name, kind=self.kind, uri=str(p), version=version, authority=authority,
            controlled=controlled, sha256=sha256_bytes(raw), size_bytes=len(raw),
            ingested_at=ingested_at or datetime.now(timezone.utc),
        )
        docx = Document(p)
        spans: list[SourceSpan] = []
        ordinal = 0
        for pi, para in enumerate(docx.paragraphs):
            text = para.text.strip()
            if text:
                spans.append(SourceSpan(
                    span_id=f"{source_id}:P{pi}", source_id=source_id, paragraph_index=pi,
                    text=text, span_hash=sha256_text(text), metadata={"style": para.style.name if para.style else None}
                ))
                ordinal += 1
        for ti, table in enumerate(docx.tables):
            for ri, row in enumerate(table.rows):
                text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if text:
                    spans.append(SourceSpan(
                        span_id=f"{source_id}:T{ti}:R{ri}", source_id=source_id,
                        text=text, span_hash=sha256_text(text), metadata={"table": ti, "row": ri}
                    ))
        return ParseResult(doc_meta, spans)


class XLSXParser(DocumentParser):
    kind = SourceKind.XLSX

    def parse(self, path, *, source_id, organization_id, audit_case_id=None, controlled=True,
              version=None, authority=None, ingested_at=None) -> ParseResult:
        from datetime import datetime, timezone
        from openpyxl import load_workbook
        p = Path(path)
        raw = p.read_bytes()
        doc_meta = SourceDocument(
            source_id=source_id, organization_id=organization_id, audit_case_id=audit_case_id,
            title=p.name, kind=self.kind, uri=str(p), version=version, authority=authority,
            controlled=controlled, sha256=sha256_bytes(raw), size_bytes=len(raw),
            ingested_at=ingested_at or datetime.now(timezone.utc),
        )
        wb = load_workbook(p, read_only=True, data_only=True)
        spans: list[SourceSpan] = []
        for ws in wb.worksheets:
            for ri, row in enumerate(ws.iter_rows(values_only=True), start=1):
                vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if not vals:
                    continue
                text = " | ".join(vals)
                spans.append(SourceSpan(
                    span_id=f"{source_id}:{ws.title}:R{ri}", source_id=source_id,
                    sheet=ws.title, cell_range=f"{ri}:{ri}", text=text,
                    span_hash=sha256_text(text), metadata={"row": ri}
                ))
        return ParseResult(doc_meta, spans)


def parser_for_path(path: str | Path) -> DocumentParser:
    ext = Path(path).suffix.lower()
    if ext == ".pdf": return PDFParser()
    if ext == ".docx": return DOCXParser()
    if ext in {".xlsx", ".xlsm"}: return XLSXParser()
    if ext in {".txt", ".md", ".csv", ".json", ".jsonl"}: return TextParser()
    raise ValueError(f"unsupported source type: {ext}")
