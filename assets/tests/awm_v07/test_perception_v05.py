from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from openpyxl import Workbook
import fitz

from aias_awm.perception import PerceptionPipeline, CandidateState, SourceKind
from aias_awm.retrieval import HybridRetriever
from aias_awm.provenance import ControlledSourceRegistry, ProvenanceVerifier
from aias_awm.perception.promotion import EvidencePromotionService
from aias_awm.domain.models import EpistemicState


def test_text_parse_chunk_candidate_and_promotion(tmp_path):
    p=tmp_path/'sample.txt'; p.write_text('Pilot has 12 errors.\nNo formal acceptance criteria were established.', encoding='utf-8')
    pipeline=PerceptionPipeline()
    result=pipeline.ingest(p, source_id='SRC-1', organization_id='ORG-1', audit_case_id='CASE-1',
        requirement_id='AR-6.1.3-E06', patterns={'AR-6.1.3-E06':['acceptance criteria','effectiveness']})
    assert result.parse.document.kind == SourceKind.TEXT
    assert result.parse.document.controlled is True
    assert result.chunks
    assert result.candidates
    promoted=pipeline.promotion_service().promote_presented(result.candidates[0])
    assert promoted.evidence is not None
    assert promoted.evidence.epistemic_state == EpistemicState.PRESENTED
    verified=pipeline.promotion_service().verify(promoted.evidence, verified_by='AUDITOR-1')
    assert verified.epistemic_state == EpistemicState.VERIFIED
    assert verified.verified_by == 'AUDITOR-1'


def test_uncontrolled_source_cannot_be_promoted(tmp_path):
    p=tmp_path/'sample.txt'; p.write_text('some evidence', encoding='utf-8')
    pipeline=PerceptionPipeline()
    result=pipeline.ingest(p, source_id='SRC-X', organization_id='ORG-1', audit_case_id='CASE-1', controlled=False)
    promoted=pipeline.promotion_service().promote_presented(result.candidates[0])
    assert promoted.evidence is None
    assert 'SOURCE_NOT_CONTROLLED' in promoted.reason


def test_pdf_docx_xlsx_parsers(tmp_path):
    pdf=tmp_path/'a.pdf'
    d=fitz.open(); pg=d.new_page(); pg.insert_text((72,72),'Effectiveness evaluation record'); d.save(pdf); d.close()
    r=PerceptionPipeline().ingest(pdf, source_id='PDF', organization_id='ORG')
    assert r.parse.document.kind == SourceKind.PDF and r.parse.spans

    docx=tmp_path/'a.docx'; doc=Document(); doc.add_paragraph('Approved audit plan'); doc.save(docx)
    r=PerceptionPipeline().ingest(docx, source_id='DOCX', organization_id='ORG')
    assert r.parse.document.kind == SourceKind.DOCX and r.parse.spans

    xlsx=tmp_path/'a.xlsx'; wb=Workbook(); ws=wb.active; ws['A1']='KPI'; ws['B1']=4.2; wb.save(xlsx)
    r=PerceptionPipeline().ingest(xlsx, source_id='XLSX', organization_id='ORG')
    assert r.parse.document.kind == SourceKind.XLSX and r.parse.spans


def test_hybrid_retrieval_is_constraint_aware(tmp_path):
    p1=tmp_path/'a.txt'; p1.write_text('formal effectiveness acceptance criteria for opportunity pilot', encoding='utf-8')
    p2=tmp_path/'b.txt'; p2.write_text('supplier calibration certificate', encoding='utf-8')
    pipe=PerceptionPipeline(); a=pipe.ingest(p1, source_id='S1', organization_id='O', audit_case_id='C'); b=pipe.ingest(p2, source_id='S2', organization_id='O', audit_case_id='C')
    ret=HybridRetriever(a.chunks+b.chunks)
    hits=ret.search('effectiveness acceptance criteria', allowed_source_ids={'S1'}, top_k=3)
    assert hits and all(h.source_id=='S1' for h in hits)
    assert hits[0].score > 0
